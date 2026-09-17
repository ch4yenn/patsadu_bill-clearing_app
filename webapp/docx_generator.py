import docx
import os
import copy
from webapp.bahttext import baht_text

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates_master")

def get_template_path(doc_type: str, approver_type: str) -> str:
    """
    doc_type: 'buy' or 'hire'
    approver_type: 'director' or 'deputy'
    """
    if doc_type == "buy":
        fname = "Form_buy_director.docx" if approver_type == "director" else "Form_buy_deputy.docx"
    else:
        fname = "Form_hire_director.docx" if approver_type == "director" else "Form_hire_deputy.docx"
    return os.path.join(TEMPLATES_DIR, fname)

def replace_text_in_paragraph(p, old_text, new_text, start_search=0):
    """
    Replaces old_text with new_text across runs in paragraph p,
    preserving font names, sizes, bolding, colors, and paragraph styles.
    Returns next search position in updated text, or -1 if no match.
    """
    if not old_text or str(old_text) == str(new_text):
        return -1
    
    full_text = "".join(r.text for r in p.runs)
    if start_search >= len(full_text):
        return -1
        
    start_idx = full_text.find(old_text, start_search)
    if start_idx == -1:
        return -1
        
    end_idx = start_idx + len(old_text)
    
    curr_len = 0
    start_run = None
    start_offset = 0
    end_run = None
    end_offset = 0
    
    for i, r in enumerate(p.runs):
        r_len = len(r.text)
        if start_run is None and curr_len + r_len > start_idx:
            start_run = i
            start_offset = start_idx - curr_len
        if end_run is None and curr_len + r_len >= end_idx:
            end_run = i
            end_offset = end_idx - curr_len
            break
        curr_len += r_len
        
    if start_run is None or end_run is None:
        return -1
        
    new_str = str(new_text)
    if start_run == end_run:
        r = p.runs[start_run]
        r.text = r.text[:start_offset] + new_str + r.text[end_offset:]
    else:
        p.runs[start_run].text = p.runs[start_run].text[:start_offset] + new_str
        for i in range(start_run + 1, end_run):
            p.runs[i].text = ""
        p.runs[end_run].text = p.runs[end_run].text[end_offset:]
        
    return start_idx + len(new_str)

def replace_in_paragraphs(paragraphs, old_text, new_text):
    if not old_text or str(old_text) == str(new_text):
        return 0
    count = 0
    old_str = str(old_text)
    new_str = str(new_text)
    for p in paragraphs:
        start_search = 0
        loop_guard = 0
        while loop_guard < 50:
            loop_guard += 1
            next_pos = replace_text_in_paragraph(p, old_str, new_str, start_search)
            if next_pos == -1:
                break
            count += 1
            start_search = max(next_pos, start_search + 1)
    return count

def format_money(val) -> str:
    try:
        f = float(val)
        return f"{f:,.2f}"
    except (ValueError, TypeError):
        return str(val)

def split_baht_satang(val) -> tuple[str, str]:
    try:
        f = float(val)
        s = f"{f:,.2f}"
        parts = s.split(".")
        return parts[0], parts[1]
    except Exception:
        return str(val), "00"

def populate_table_0(table, items, total_amount, baht_text_str, is_hire=False):
    """
    Table 0 (รายละเอียดแนบท้าย):
    Header: Rows 0 and 1
    Item rows: row 2 onwards
    Last row: Total row
    """
    num_items = len(items)
    total_rows = len(table.rows)
    total_row_idx = total_rows - 1
    
    # Ensure header Col 1 has regular font and no Heading 1 style
    from docx.oxml.ns import qn
    for r_idx in [0, 1]:
        if len(table.rows) > r_idx and len(table.rows[r_idx].cells) > 1:
            h_cell = table.rows[r_idx].cells[1]
            for p in h_cell.paragraphs:
                pPr = p._p.get_or_add_pPr()
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None:
                    pPr.remove(pStyle)
                p_rPr = pPr.find(qn('w:rPr'))
                if p_rPr is not None:
                    for b_tag in [qn('w:b'), qn('w:bCs')]:
                        b = p_rPr.find(b_tag)
                        if b is not None and b.attrib.get(qn('w:val')) != '0':
                            p_rPr.remove(b)
                p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.bold = False
                    rPr = run._r.get_or_add_rPr()
                    for b_tag in [qn('w:b'), qn('w:bCs')]:
                        b = rPr.find(b_tag)
                        if b is not None and b.attrib.get(qn('w:val')) != '0':
                            rPr.remove(b)
    
    template_item_rows_count = total_rows - 3 # row 0, 1 (headers), last row (total)
    
    # Trim excess dummy rows
    if num_items < template_item_rows_count:
        for r_idx in range(total_row_idx - 1, 1 + num_items, -1):
            tr = table.rows[r_idx]._tr
            tr.getparent().remove(tr)
    # Clone row 2 if items exceed template rows
    elif num_items > template_item_rows_count:
        sample_tr = table.rows[2]._tr
        last_tr = table.rows[-1]._tr
        for _ in range(num_items - template_item_rows_count):
            new_tr = copy.deepcopy(sample_tr)
            last_tr.addprevious(new_tr)
            
    # Populate item rows
    for idx, item in enumerate(items):
        r_idx = 2 + idx
        row = table.rows[r_idx]
        qty_unit = f"{item.get('qty', '')} {item.get('unit', '')}".strip()
        price_b, price_s = split_baht_satang(item.get('price_per_unit', 0))
        tot_b, tot_s = split_baht_satang(item.get('total_price', 0))
        
        # Col 0: ลำดับ
        if len(row.cells) > 0:
            row.cells[0].text = str(idx + 1)
        # Col 1: รายละเอียด
        if len(row.cells) > 1:
            row.cells[1].text = str(item.get('name', ''))
        # Col 2: จำนวน หน่วย
        if len(row.cells) > 2:
            row.cells[2].text = qty_unit
        # Col 3-4: ราคามาตรฐาน (บาท, สต.)
        if len(row.cells) > 3:
            row.cells[3].text = price_b
        if len(row.cells) > 4:
            row.cells[4].text = price_s
        # Col 5-6: หน่วยละ (บาท, สต.)
        if len(row.cells) > 5:
            row.cells[5].text = price_b
        if len(row.cells) > 6:
            row.cells[6].text = price_s
        # Col 7-8: จำนวนเงิน (บาท, สต.)
        if len(row.cells) > 7:
            row.cells[7].text = tot_b
        if len(row.cells) > 8:
            row.cells[8].text = tot_s
            
        # Ensure alignments
        if len(row.cells) > 0 and row.cells[0].paragraphs:
            row.cells[0].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        if len(row.cells) > 1 and row.cells[1].paragraphs:
            row.cells[1].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
        if len(row.cells) > 2 and row.cells[2].paragraphs:
            row.cells[2].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        for c in range(3, min(9, len(row.cells))):
            if row.cells[c].paragraphs:
                row.cells[c].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT

    # Update Total row (the last row)
    last_row = table.rows[-1]
    tot_b, tot_s = split_baht_satang(total_amount)
    if len(last_row.cells) > 0:
        last_row.cells[0].text = f"รวมเป็นเงินทั้งสิ้น  ({baht_text_str})"
        if last_row.cells[0].paragraphs:
            last_row.cells[0].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    if len(last_row.cells) > 7:
        last_row.cells[7].text = tot_b
        if last_row.cells[7].paragraphs:
            last_row.cells[7].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
    if len(last_row.cells) > 8:
        last_row.cells[8].text = tot_s
        if last_row.cells[8].paragraphs:
            last_row.cells[8].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT

def populate_table_1(table, items, financial, baht_text_str):
    """
    Table 1 (ใบสั่งซื้อ/ใบสั่งจ้าง):
    Header: Row 0
    Item rows: row 1 to row len-4
    Summary rows: last 3 rows (รวมเป็นเงิน, ภาษีมูลค่าเพิ่ม, รวมเป็นเงินทั้งสิ้น)
    """
    num_items = len(items)
    total_rows = len(table.rows)
    template_item_rows_count = total_rows - 4 # row 0 (header) + 3 summary rows
    
    # Trimming extra rows
    if num_items < template_item_rows_count:
        for r_idx in range(total_rows - 4, num_items, -1):
            tr = table.rows[r_idx]._tr
            tr.getparent().remove(tr)
    elif num_items > template_item_rows_count:
        sample_tr = table.rows[1]._tr
        first_summary_tr = table.rows[-3]._tr
        for _ in range(num_items - template_item_rows_count):
            new_tr = copy.deepcopy(sample_tr)
            first_summary_tr.addprevious(new_tr)
            
    # Populate items
    for idx, item in enumerate(items):
        r_idx = 1 + idx
        row = table.rows[r_idx]
        p_unit_str = format_money(item.get('price_per_unit', 0))
        tot_str = format_money(item.get('total_price', 0))
        
        # Col 0: ลำดับ
        if len(row.cells) > 0:
            row.cells[0].text = str(idx + 1)
        # Col 1: รายการ
        if len(row.cells) > 1:
            row.cells[1].text = str(item.get('name', ''))
        # Col 2: จำนวน
        if len(row.cells) > 2:
            row.cells[2].text = str(item.get('qty', ''))
        # Col 3: หน่วย
        if len(row.cells) > 3:
            row.cells[3].text = str(item.get('unit', ''))
        # Col 4: ราคาต่อหน่วย (บาท)
        if len(row.cells) > 4:
            row.cells[4].text = f" {p_unit_str} "
        # Col 5: จำนวนเงิน (บาท)
        if len(row.cells) > 5:
            row.cells[5].text = f" {tot_str} "
            
        # Alignments
        if len(row.cells) > 0 and row.cells[0].paragraphs:
            row.cells[0].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        if len(row.cells) > 1 and row.cells[1].paragraphs:
            row.cells[1].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
        if len(row.cells) > 2 and row.cells[2].paragraphs:
            row.cells[2].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        if len(row.cells) > 3 and row.cells[3].paragraphs:
            row.cells[3].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        if len(row.cells) > 4 and row.cells[4].paragraphs:
            row.cells[4].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
        if len(row.cells) > 5 and row.cells[5].paragraphs:
            row.cells[5].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT

    # Summary rows (last 3 rows):
    subtotal_row = table.rows[-3]
    vat_row = table.rows[-2]
    total_row = table.rows[-1]
    
    # Left merged cell gets the BahtText: (คำอ่าน)
    for r in [subtotal_row, vat_row, total_row]:
        if len(r.cells) > 0:
            r.cells[0].text = f"({baht_text_str})"
            if r.cells[0].paragraphs:
                r.cells[0].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER

    subtotal_val = format_money(financial.get('goods_value', financial.get('total_amount', 0)))
    vat_val = format_money(financial.get('vat', 0)) if financial.get('has_vat') else "-"
    total_val = format_money(financial.get('total_amount', 0))

    if len(subtotal_row.cells) > 5:
        subtotal_row.cells[5].text = f" {subtotal_val} "
        if subtotal_row.cells[5].paragraphs:
            subtotal_row.cells[5].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT

    if len(vat_row.cells) > 5:
        vat_row.cells[5].text = f" {vat_val} "
        if vat_row.cells[5].paragraphs:
            vat_row.cells[5].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT

    if len(total_row.cells) > 5:
        total_row.cells[5].text = f" {total_val} "
        if total_row.cells[5].paragraphs:
            total_row.cells[5].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT

def generate_procurement_doc(data: dict, output_path: str) -> str:
    """
    Main entry point to generate the 6-page procurement package Word document.
    Uses section-scoped replacement to guarantee zero cascading text conflicts.
    """
    doc_type = data.get("doc_type", "buy") # 'buy' or 'hire'
    approver_type = data.get("approver_type", "director") # 'director' or 'deputy'
    
    template_file = get_template_path(doc_type, approver_type)
    if not os.path.exists(template_file):
        raise FileNotFoundError(f"Template not found: {template_file}")
        
    doc = docx.Document(template_file)
    
    # Extract data fields
    dept = str(data.get("department", "กลุ่มบริหารทั่วไป")).strip()
    project_name = str(data.get("project_name", "")).strip()
    hire_job_name = str(data.get("hire_job_name", "ทำป้ายไวนิล")).strip()
    items = data.get("items", [])
    num_items = len(items)
    
    # Financial calculation
    total_amount = sum(float(item.get("total_price", 0)) for item in items)
    has_vat = data.get("has_vat", False)
    if has_vat:
        goods_value = round(total_amount * 100 / 107, 2)
        vat = round(total_amount - goods_value, 2)
    else:
        goods_value = total_amount
        vat = 0.0
        
    has_wht = data.get("has_withholding_tax", False)
    wht_rate = float(data.get("withholding_rate", 1.0)) if has_wht else 0.0
    withholding_tax = round(goods_value * (wht_rate / 100), 2) if has_wht else 0.0
    fine = float(data.get("fine_amount", 0.0)) if data.get("fine_amount") else 0.0
    net_pay = round(total_amount - withholding_tax - fine, 2)
    
    baht_text_total = baht_text(total_amount)
    baht_text_net = baht_text(net_pay)
    
    financial = {
        "goods_value": goods_value,
        "has_vat": has_vat,
        "vat": vat,
        "total_amount": total_amount,
        "has_withholding_tax": has_wht,
        "withholding_tax": withholding_tax,
        "fine": fine,
        "net_pay": net_pay
    }
    
    # Dates and Numbers
    doc1_no = str(data.get("doc1_no", "")).strip()
    doc1_date = str(data.get("doc1_date", "")).strip()
    doc2_no = str(data.get("doc2_no", "")).strip()
    doc2_date = str(data.get("doc2_date", "")).strip()
    doc3_no = str(data.get("doc3_no", "")).strip()
    doc3_date = str(data.get("doc3_date", "")).strip()
    
    po_no = str(data.get("po_no", "")).strip()
    po_date = str(data.get("po_date", "")).strip()
    delivery_days = str(data.get("delivery_days", 5)).strip()
    delivery_due_date = str(data.get("delivery_due_date", "")).strip()
    delivery_actual_date = str(data.get("delivery_actual_date", "")).strip()
    
    # Vendor
    vendor = data.get("vendor", {})
    vendor_name = str(vendor.get("name", "")).strip()
    vendor_tax = str(vendor.get("tax_id", "")).strip()
    vendor_addr = str(vendor.get("address", "")).strip()
    vendor_subdist = str(vendor.get("subdistrict", "")).strip()
    vendor_dist = str(vendor.get("district", "")).strip()
    vendor_prov = str(vendor.get("province", "")).strip()
    vendor_zip = str(vendor.get("zipcode", "")).strip()
    vendor_full_location = f"ตำบล{vendor_subdist} อำเภอ{vendor_dist} จังหวัด{vendor_prov} {vendor_zip}".strip()
    vendor_phone = str(vendor.get("phone", "-")).strip()
    vendor_signer = str(vendor.get("signer_name", "")).strip()
    vendor_signer_pos = str(vendor.get("signer_position", "เจ้าของกิจการ")).strip()
    
    # Receipt
    receipt = data.get("receipt", {})
    receipt_type = str(receipt.get("type", "ใบเสร็จรับเงิน")).strip()
    receipt_book = str(receipt.get("book_no", "-")).strip()
    receipt_no = str(receipt.get("no", "-")).strip()
    receipt_date = str(receipt.get("date", delivery_actual_date)).strip()
    
    # Committee
    committee = data.get("committee", [
        {"name": "นางสาวธัญญาภรณ์  สุกันทา", "position": "ครู"},
        {"name": "นางสาวกนกวรรณ  มีเทียม", "position": "ครู"},
        {"name": "นายภูวดล  บุญต่อ", "position": "ครูผู้ช่วย"}
    ])
    while len(committee) < 3:
        committee.append({"name": "-", "position": "-"})
        
    # Officers
    officer_supplies = data.get("officer_supplies", "นายภาณุพงศ์  อัมพรภาค")
    head_supplies = data.get("head_supplies", "อดิศักดิ์  บัวดี" if approver_type == "director" else "นางสาวจันจิรา  น่วมนวล")
    finance_officer = data.get("finance_officer", "นางสาวอังคณา  เสถียรดี")
    
    # -------------------------------------------------------------
    # 1. Update Tables (Table 0 and Table 1)
    # -------------------------------------------------------------
    if len(doc.tables) >= 2:
        populate_table_0(doc.tables[0], items, total_amount, baht_text_total, is_hire=(doc_type == "hire"))
        populate_table_1(doc.tables[1], items, financial, baht_text_total)

    # -------------------------------------------------------------
    # 2. Identify paragraph boundaries of the 6 documents
    # -------------------------------------------------------------
    p_texts = [p.text.strip() for p in doc.paragraphs]
    
    idx_p1 = 0
    idx_p2 = None
    idx_p3 = None
    idx_p4 = None
    idx_p5 = None
    idx_p6 = None
    
    for i, t in enumerate(p_texts):
        if "รายละเอียดแนบท้ายบันทึกข้อความ" in t and idx_p2 is None:
            idx_p2 = i
        elif "บันทึกข้อความ" in t:
            if idx_p2 is not None and idx_p3 is None and i > idx_p2:
                idx_p3 = i
            elif idx_p5 is not None and idx_p6 is None and i > idx_p5:
                idx_p6 = i
        elif ("ใบสั่งซื้อ" in t or "ใบสั่งจ้าง" in t) and idx_p4 is None:
            if idx_p3 is not None and i > idx_p3:
                idx_p4 = i
        elif "ใบตรวจรับพัสดุ" in t and idx_p5 is None:
            if idx_p4 is not None and i > idx_p4:
                idx_p5 = i
                
    # Fallback bounds if any marker shifted
    if idx_p2 is None: idx_p2 = 31
    if idx_p3 is None: idx_p3 = 43
    if idx_p4 is None: idx_p4 = 75
    if idx_p5 is None: idx_p5 = 96
    if idx_p6 is None: idx_p6 = 119
    
    sec1_paras = doc.paragraphs[idx_p1:idx_p2]
    sec2_paras = doc.paragraphs[idx_p2:idx_p3]
    sec3_paras = doc.paragraphs[idx_p3:idx_p4]
    sec4_paras = doc.paragraphs[idx_p4:idx_p5]
    sec5_paras = doc.paragraphs[idx_p5:idx_p6]
    sec6_paras = doc.paragraphs[idx_p6:]
    
    # -------------------------------------------------------------
    # Template Original Values Mapping
    # -------------------------------------------------------------
    if doc_type == "buy" and approver_type == "director":
        orig_dept = "ฝ่ายบริหารงานหน่วยบริการ"
        orig_reason = "กิจกรรมชุมนุม"
        orig_items_count = "6"
        orig_amt = "500.00"
        orig_amt_text = "ห้าร้อยบาทถ้วน"
        orig_days = "5"
        orig_doc1_no = "2912"
        orig_doc1_date = "30 กรกฎาคม 2569"
        orig_doc2_no = "2913"
        orig_doc2_date = "31 กรกฎาคม 2569"
        orig_po_no = "91/2569"
        orig_po_date = "31 กรกฎาคม 2569"
        orig_due_date = "5 สิงหาคม 2569"
        orig_vendor = "ห้างหุ้นส่วนจำกัด น้ำฝน"
        orig_tax = "0643567000041"
        orig_addr = "เลขที่ 334 หมู่ 8"
        orig_loc = "ตำบลทุ่งเสลี่ยม อำเภอทุ่งเสลี่ยม จังหวัดสุโขทัย 64150"
        orig_phone = "082-0071206"
        orig_signer = "นางสาวแพรพลอย  ขวัญนาค"
        orig_signer_pos = "เจ้าของกิจการ"
        orig_delivery_date = "5 สิงหาคม 2569"
        orig_rec_type = "ใบเสร็จรับเงิน"
        orig_rec_book = "5"
        orig_rec_no = "230"
        orig_rec_date = "5 สิงหาคม 2569"
        orig_doc3_no = "2914"
        orig_doc3_date = "10 สิงหาคม 2569"
        orig_c1 = "นางสาวธัญญาภรณ์  สุกันทา"
        orig_c1_pos = "ครู"
        orig_c2 = "นางสาวกนกวรรณ  มีเทียม"
        orig_c2_pos = "ครู "
        orig_c3 = "นายภูวดล  บุญต่อ"
        orig_c3_pos = "ครูผู้ช่วย"
        
    elif doc_type == "buy" and approver_type == "deputy":
        orig_dept = "ฝ่ายบริหารทั่วไป"
        orig_reason = "ให้รถยนต์ส่วนกลาง ทะเบียน กฉ 6739 สุโขทัย มีความคุ้มครองและสามารถใช้งานในราชการได้อย่างถูกต้องตามกฎหมาย"
        orig_items_count = "1"
        orig_amt = "645.21"
        orig_amt_text = "หกร้อยสี่สิบห้าบาทยี่สิบเอ็ดสตางค์"
        orig_days = "3"
        orig_doc1_no = "2063"
        orig_doc1_date = "16 เมษายน 2569"
        orig_doc2_no = "2064"
        orig_doc2_date = "17 เมษายน 2569"
        orig_po_no = "59/2569"
        orig_po_date = "17 เมษายน 2569"
        orig_due_date = "20 เมษายน 2569"
        orig_vendor = "บริษัท สหมงคลประกันภัย จำกัด (มหาชน) (สำนักงานใหญ่)"
        orig_tax = "0107555000333"
        orig_addr = "เลขที่ 7 ซอยสาทร 11 หมู่ -"
        orig_loc = "ตำบลแขวงยานนาวา อำเภอเขตสาทร จังหวัดกรุงเทพฯ 10120"
        orig_phone = "0-26877777"
        orig_signer = "ปวิรตา  เศรษฐบุตร"
        orig_signer_pos = "เจ้าหน้าที่การเงิน"
        orig_delivery_date = "20 เมษายน 2569"
        orig_rec_type = "หนังสือส่งมอบของบริษัท สหมงคลประกันภัย จำกัด (มหาชน) (สำนักงานใหญ่)"
        orig_rec_book = "-"
        orig_rec_no = "C187617"
        orig_rec_date = "20 เมษายน 2569"
        orig_doc3_no = "2065"
        orig_doc3_date = "21 เมษายน 2569"
        orig_c1 = "นายอานนท์  เขียวแก้ว"
        orig_c1_pos = "ครู"
        orig_c2 = "ว่าที่ร.ต.กำพล  ไชยสวัสดิ์"
        orig_c2_pos = "ครูชำนาญการ"
        orig_c3 = "นายทินธารัตน์  แก้วมุกดา"
        orig_c3_pos = "ครูชำนาญการ"
        
    elif doc_type == "hire" and approver_type == "director":
        orig_dept = "หน่วยบริการศรีสำโรง"
        orig_reason = 'โครงการศิลปะบำบัดสำหรับเด็ก "Art Hug กอดรักด้วยศิลปะ"'
        orig_items_count = "1"
        orig_amt = "450.00"
        orig_amt_text = "สี่ร้อยห้าสิบบาทถ้วน"
        orig_days = "3"
        orig_doc1_no = "2724"
        orig_doc1_date = "16 กรกฎาคม 2569"
        orig_doc2_no = "2725"
        orig_doc2_date = "17 กรกฎาคม 2569"
        orig_po_no = "592/2569"
        orig_po_date = "17 กรกฎาคม 2569"
        orig_due_date = "20 กรกฎาคม 2569"
        orig_vendor = "ร้านอีสแอนด์ดี ดีไซน์"
        orig_tax = "1640100071271"
        orig_addr = "เลขที่ 27/3 หมู่ 2"
        orig_loc = "ตำบลปากพระ อำเภอเมืองสุโขทัย จังหวัดสุโขทัย 64000"
        orig_phone = "095-6639228"
        orig_signer = "นายธนพล  กันเหม็น"
        orig_signer_pos = "ผู้รับจ้าง"
        orig_delivery_date = "20 กรกฎาคม 2569"
        orig_rec_type = "ใบสำคัญรับเงินค่าทำป้ายไวนิล"
        orig_rec_book = "57"
        orig_rec_no = "2856"
        orig_rec_date = "20 กรกฎาคม 2569"
        orig_doc3_no = "2726"
        orig_doc3_date = "22 กรกฎาคม 2569"
        orig_c1 = "นางสาวชนิดา  แก้วนาค"
        orig_c1_pos = "ครู"
        orig_c2 = "นางวัชราภรณ์ ปะสะจัน"
        orig_c2_pos = "ครูชำนาญการ"
        orig_c3 = "นางสาวระพีพรรณ บัวผัน"
        orig_c3_pos = "ครูชำนาญการ"
        
    else: # hire and deputy
        orig_dept = "หน่วยบริการบ้านด่านลานหอย"
        orig_reason = "โครงการธาราบำบัดเพื่อส่งเสริมสุขภาพและกระตุ้นพัฒนาการเด็กพิการ ประจำปีงบประมาณ พ.ศ.2569"
        orig_items_count = "1"
        orig_amt = "450.00"
        orig_amt_text = "สี่ร้อยห้าสิบบาทถ้วน"
        orig_days = "5"
        orig_doc1_no = "2861"
        orig_doc1_date = "21 พฤษภาคม 2569"
        orig_doc2_no = "2862"
        orig_doc2_date = "22 พฤษภาคม 2569"
        orig_po_no = "620/2569"
        orig_po_date = "22 พฤษภาคม 2569"
        orig_due_date = "27 พฤษภาคม 2569"
        orig_vendor = "ร้านอีสแอนด์ดี ดีไซน์"
        orig_tax = "1640100071271"
        orig_addr = "เลขที่ 27/3 หมู่ 2"
        orig_loc = "ตำบลปากพระ อำเภอเมืองสุโขทัย จังหวัดสุโขทัย 64000"
        orig_phone = "095-6639228"
        orig_signer = "นายธนพล  กันเหม็น"
        orig_signer_pos = "ผู้รับจ้าง"
        orig_delivery_date = "26 พฤษภาคม 2569"
        orig_rec_type = "ใบสำคัญรับเงินค่าทำป้ายไวนิลโครงการ"
        orig_rec_book = "56"
        orig_rec_no = "2763"
        orig_rec_date = "26 พฤษภาคม 2569"
        orig_doc3_no = "2863"
        orig_doc3_date = " กันยายน 2569"
        orig_c1 = "นางสาวอรณี ทองปอด"
        orig_c1_pos = "ครู"
        orig_c2 = "นางสาวปรียากร สอนโพธิ์"
        orig_c2_pos = "ครูชำนาญการ"
        orig_c3 = "นางวัชราภรณ์ ปะสะจัน"
        orig_c3_pos = "ครูชำนาญการ"

    # Helper lambda to apply standard amount replacements
    def apply_amount_replaces(paras):
        replace_in_paragraphs(paras, orig_amt_text, baht_text_total)
        replace_in_paragraphs(paras, f"  {orig_amt} บาท", f"  {format_money(total_amount)} บาท")
        replace_in_paragraphs(paras, f" {orig_amt} บาท", f" {format_money(total_amount)} บาท")
        replace_in_paragraphs(paras, f"{orig_amt} บาท", f"{format_money(total_amount)} บาท")
        replace_in_paragraphs(paras, f" {orig_amt}บาท", f" {format_money(total_amount)} บาท")

    # Helper lambda for items count
    def apply_item_count_replaces(paras):
        replace_in_paragraphs(paras, f"จำนวน {orig_items_count} รายการ", f"จำนวน {num_items} รายการ")
        replace_in_paragraphs(paras, f"จำนวน      {orig_items_count}       รายการ", f"จำนวน      {num_items}       รายการ")

    # -------------------------------------------------------------
    # SECTION 1: รายงานขอซื้อ / ขอจ้าง
    # -------------------------------------------------------------
    if doc_type == "hire":
        replace_in_paragraphs(sec1_paras, "ทำป้ายไวนิลโครงการ", hire_job_name)
        replace_in_paragraphs(sec1_paras, "ทำป้ายไวนิล", hire_job_name)
    replace_in_paragraphs(sec1_paras, orig_dept, dept)
    replace_in_paragraphs(sec1_paras, orig_reason, project_name)
    apply_item_count_replaces(sec1_paras)
    apply_amount_replaces(sec1_paras)
    replace_in_paragraphs(sec1_paras, f"/{orig_doc1_no}", f"/{doc1_no}")
    replace_in_paragraphs(sec1_paras, orig_doc1_date, doc1_date)
    replace_in_paragraphs(sec1_paras, f"ภายใน {orig_days} วัน", f"ภายใน {delivery_days} วัน")

    # Committee strictly within their respective paragraphs (prevents position swapping)
    p_c1 = None
    p_c2 = None
    p_c3 = None
    for p in sec1_paras:
        t = p.text.strip()
        if t.startswith("(1)") or "(1)" in t:
            p_c1 = p
        elif t.startswith("(2)") or "(2)" in t:
            p_c2 = p
        elif t.startswith("(3)") or "(3)" in t:
            p_c3 = p

    if p_c1:
        replace_text_in_paragraph(p_c1, orig_c1, committee[0].get("name", ""))
        replace_text_in_paragraph(p_c1, f"ตำแหน่ง {orig_c1_pos}", f"ตำแหน่ง {committee[0].get('position', '')}")
        replace_text_in_paragraph(p_c1, f"ตำแหน่ง  {orig_c1_pos}", f"ตำแหน่ง {committee[0].get('position', '')}")
    if p_c2:
        replace_text_in_paragraph(p_c2, orig_c2, committee[1].get("name", ""))
        replace_text_in_paragraph(p_c2, f"ตำแหน่ง {orig_c2_pos}", f"ตำแหน่ง {committee[1].get('position', '')}")
        replace_text_in_paragraph(p_c2, f"ตำแหน่ง  {orig_c2_pos}", f"ตำแหน่ง {committee[1].get('position', '')}")
    if p_c3:
        replace_text_in_paragraph(p_c3, orig_c3, committee[2].get("name", ""))
        replace_text_in_paragraph(p_c3, f"ตำแหน่ง {orig_c3_pos}", f"ตำแหน่ง {committee[2].get('position', '')}")
        replace_text_in_paragraph(p_c3, f"ตำแหน่ง  {orig_c3_pos}", f"ตำแหน่ง {committee[2].get('position', '')}")

    # -------------------------------------------------------------
    # SECTION 2: รายละเอียดแนบท้ายบันทึกข้อความ
    # -------------------------------------------------------------
    replace_in_paragraphs(sec2_paras, orig_dept, dept)
    apply_item_count_replaces(sec2_paras)
    replace_in_paragraphs(sec2_paras, f"ที่ {orig_doc1_no}/", f"ที่ {doc1_no}/")
    replace_in_paragraphs(sec2_paras, orig_doc1_date, doc1_date)

    # -------------------------------------------------------------
    # SECTION 3: รายงานผลการพิจารณาและอนุมัติสั่งซื้อ / สั่งจ้าง
    # -------------------------------------------------------------
    if doc_type == "hire":
        replace_in_paragraphs(sec3_paras, "ทำป้ายไวนิลโครงการ", hire_job_name)
        replace_in_paragraphs(sec3_paras, "ทำป้ายไวนิล", hire_job_name)
    replace_in_paragraphs(sec3_paras, orig_reason, project_name)
    apply_item_count_replaces(sec3_paras)
    apply_amount_replaces(sec3_paras)
    replace_in_paragraphs(sec3_paras, f"/{orig_doc2_no}", f"/{doc2_no}")
    replace_in_paragraphs(sec3_paras, orig_doc2_date, doc2_date)
    replace_in_paragraphs(sec3_paras, orig_vendor, vendor_name)
    replace_in_paragraphs(sec3_paras, f"ส่งมอบ {orig_days} วัน", f"ส่งมอบ {delivery_days} วัน")

    # -------------------------------------------------------------
    # SECTION 4: ใบสั่งซื้อ / ใบสั่งจ้าง
    # -------------------------------------------------------------
    replace_in_paragraphs(sec4_paras, orig_vendor, vendor_name)
    replace_in_paragraphs(sec4_paras, orig_tax, vendor_tax)
    replace_in_paragraphs(sec4_paras, orig_addr, vendor_addr)
    replace_in_paragraphs(sec4_paras, orig_loc, vendor_full_location)
    replace_in_paragraphs(sec4_paras, orig_phone, vendor_phone)
    replace_in_paragraphs(sec4_paras, orig_po_no, po_no)
    replace_in_paragraphs(sec4_paras, orig_po_date, po_date)
    replace_in_paragraphs(sec4_paras, f"ภายใน {orig_days} วัน", f"ภายใน {delivery_days} วัน")
    replace_in_paragraphs(sec4_paras, orig_due_date, delivery_due_date)
    replace_in_paragraphs(sec4_paras, orig_signer, vendor_signer)
    replace_in_paragraphs(sec4_paras, orig_signer_pos, vendor_signer_pos)

    # -------------------------------------------------------------
    # SECTION 5: ใบตรวจรับพัสดุ
    # -------------------------------------------------------------
    if doc_type == "hire":
        replace_in_paragraphs(sec5_paras, "ทำป้ายไวนิลโครงการ", hire_job_name)
        replace_in_paragraphs(sec5_paras, "ทำป้ายไวนิล", hire_job_name)
    replace_in_paragraphs(sec5_paras, orig_vendor, vendor_name)
    replace_in_paragraphs(sec5_paras, orig_reason, project_name)
    apply_item_count_replaces(sec5_paras)
    apply_amount_replaces(sec5_paras)
    replace_in_paragraphs(sec5_paras, orig_po_no, po_no)
    replace_in_paragraphs(sec5_paras, orig_po_date, po_date)
    replace_in_paragraphs(sec5_paras, orig_due_date, delivery_due_date)
    replace_in_paragraphs(sec5_paras, orig_delivery_date, delivery_actual_date)
    replace_in_paragraphs(sec5_paras, orig_rec_type, receipt_type)
    replace_in_paragraphs(sec5_paras, f"เล่มที่ {orig_rec_book}", f"เล่มที่ {receipt_book}")
    replace_in_paragraphs(sec5_paras, f"เลขที่ {orig_rec_no}", f"เลขที่ {receipt_no}")
    replace_in_paragraphs(sec5_paras, orig_rec_date, receipt_date)
    # Committee Signers strictly matched by paragraph
    for p in sec5_paras:
        if orig_c1 in p.text:
            replace_text_in_paragraph(p, orig_c1, committee[0].get("name", ""))
        elif orig_c2 in p.text:
            replace_text_in_paragraph(p, orig_c2, committee[1].get("name", ""))
        elif orig_c3 in p.text:
            replace_text_in_paragraph(p, orig_c3, committee[2].get("name", ""))

    # -------------------------------------------------------------
    # SECTION 6: ขออนุมัติจ่ายเงิน
    # -------------------------------------------------------------
    if doc_type == "hire":
        replace_in_paragraphs(sec6_paras, "ทำป้ายไวนิลโครงการ", hire_job_name)
        replace_in_paragraphs(sec6_paras, "ทำป้ายไวนิล", hire_job_name)
    replace_in_paragraphs(sec6_paras, orig_reason, project_name)
    apply_item_count_replaces(sec6_paras)
    apply_amount_replaces(sec6_paras)
    replace_in_paragraphs(sec6_paras, f"/{orig_doc3_no}", f"/{doc3_no}")
    replace_in_paragraphs(sec6_paras, orig_doc3_date, doc3_date)
    replace_in_paragraphs(sec6_paras, orig_delivery_date, delivery_actual_date)
    
    # Financial breakdown lines in Sec 6
    goods_str = format_money(goods_value)
    vat_str = format_money(vat) if has_vat else "-"
    req_str = format_money(total_amount)
    wht_str = format_money(withholding_tax) if has_wht else "-"
    fine_str = format_money(fine) if fine > 0 else "-"
    net_str = format_money(net_pay)
    
    for p in sec6_paras:
        pt = p.text
        if "มูลค่าสินค้าบริการ" in pt:
            p.text = f"\tมูลค่าสินค้าบริการ\t{goods_str}\tบาท"
        elif "บวกภาษีมูลค่าเพิ่ม" in pt:
            p.text = f"\tบวกภาษีมูลค่าเพิ่ม\t{vat_str}\tบาท"
        elif "จำนวนเงินที่ขอเบิก" in pt:
            p.text = f"\tจำนวนเงินที่ขอเบิก\t{req_str}\tบาท"
        elif "หักภาษีเงินได้" in pt:
            p.text = f"\tหักภาษีเงินได้\t{wht_str}\tบาท"
        elif pt.strip().startswith("ค่าปรับ") or "\tค่าปรับ\t" in pt:
            p.text = f"\tค่าปรับ\t{fine_str}\tบาท"
        elif "คงเหลือจ่ายจริง" in pt:
            p.text = f"\tคงเหลือจ่ายจริง\t{net_str}\tบาท"
        elif pt.strip() == f"({orig_amt_text})" or (pt.strip().startswith("(") and pt.strip().endswith("บาทถ้วน)")):
            p.text = f"\t\t({baht_text_net})"
            
    # Ensure mailMerge is completely stripped from settings
    for child in list(doc.settings.element):
        if child.tag.endswith("mailMerge"):
            doc.settings.element.remove(child)

    # Save output docx
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path
