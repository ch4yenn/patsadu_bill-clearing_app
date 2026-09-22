import os
import sys
import copy
import re
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn

# Dynamic path resolution (supports Windows local and Linux cloud like Render)
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PARENT_DIR = os.path.dirname(CURRENT_DIR)

for p in [PARENT_DIR, CURRENT_DIR]:
    if p not in sys.path:
        sys.path.insert(0, p)

from webapp.bahttext import baht_text

# Candidate directories for templates_master
CANDIDATE_DIRS = [
    os.path.join(PARENT_DIR, "templates_master"),
    os.path.join(CURRENT_DIR, "templates_master"),
    os.path.join(os.getcwd(), "templates_master"),
    os.path.join(os.getcwd(), "procurement-web-app", "templates_master"),
]

FONT_NAME = "TH SarabunIT๙"

def get_template_path(doc_type: str, approver_type: str) -> str:
    if doc_type == "buy":
        fname = "Form_buy_director.docx" if approver_type == "director" else "Form_buy_deputy.docx"
    else:
        fname = "Form_hire_director.docx" if approver_type == "director" else "Form_hire_deputy.docx"
        
    for d in CANDIDATE_DIRS:
        candidate = os.path.join(d, fname)
        if os.path.exists(candidate):
            return candidate
            
    return os.path.join(PARENT_DIR, "templates_master", fname)

def format_money(val) -> str:
    try:
        f = float(val)
        return f"{f:,.2f}"
    except (ValueError, TypeError):
        return str(val)

def split_baht_satang(val) -> tuple[str, str]:
    try:
        f = float(val)
        s = f"{f:,.2f}"
        parts = s.split(".")
        return parts[0], parts[1]
    except Exception:
        return str(val), "00"

def parse_vendor_address(addr_str: str) -> tuple[str, str]:
    s = str(addr_str or "").strip()
    s = re.sub(r'^(?:บ้าน)?เลขที่\s*', '', s)
    match = re.search(r'(?:หมู่ที่|หมู่|ม\.)\s*([0-9\u0E50-\u0E59]+)', s)
    if match:
        moo = match.group(1)
        house_no = re.sub(r'(?:หมู่ที่|หมู่|ม\.)\s*[0-9\u0E50-\u0E59]+', '', s).strip()
        return house_no or s, moo
    return s, "-"

def extract_thai_year(date_str: str) -> str:
    match = re.search(r'(25[0-9]{2}|๒๕[๐-๙]{2})', str(date_str))
    if match:
        return match.group(1)
    return "2569"

def clean_doc_no(val: str) -> str:
    s = str(val or "").strip()
    s = re.sub(r'^(?:ศธ\.?\s*[๐0][๔4][๐0][๐0][๗7]\.[๗7][๔4][๐0]/|ศธ\s*[0-9/.\s]+/?)', '', s).strip()
    return s or val

def clean_officer_name(s: str) -> str:
    s = str(s or "").strip()
    if (s.startswith("(") and s.endswith(")")) or (s.startswith("（") and s.endswith("）")):
        s = s[1:-1].strip()
    return s

def replace_text_in_paragraph(p, old_text, new_text, start_search=0):
    if not old_text or str(old_text) == str(new_text):
        return -1
    full_text = "".join(r.text for r in p.runs)
    if start_search >= len(full_text):
        return -1
    start_idx = full_text.find(old_text, start_search)
    if start_idx == -1:
        return -1
    end_idx = start_idx + len(old_text)
    
    curr_len = 0
    start_run = None
    start_offset = 0
    end_run = None
    end_offset = 0
    
    for i, r in enumerate(p.runs):
        r_len = len(r.text)
        if start_run is None and curr_len + r_len > start_idx:
            start_run = i
            start_offset = start_idx - curr_len
        if end_run is None and curr_len + r_len >= end_idx:
            end_run = i
            end_offset = end_idx - curr_len
            break
        curr_len += r_len
        
    if start_run is None or end_run is None:
        return -1
        
    new_str = str(new_text)
    if start_run == end_run:
        r = p.runs[start_run]
        r.text = r.text[:start_offset] + new_str + r.text[end_offset:]
    else:
        p.runs[start_run].text = p.runs[start_run].text[:start_offset] + new_str
        for i in range(start_run + 1, end_run):
            p.runs[i].text = ""
        p.runs[end_run].text = p.runs[end_run].text[end_offset:]
        
    return start_idx + len(new_str)

def set_cell_formatted_text(cell, text, align=None, bold=False, font_name=FONT_NAME, font_size=16):
    cell.text = str(text)
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    for r in p.runs:
        r.bold = bold
        r.font.name = font_name
        r.font.size = Pt(font_size)
        rPr = r._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = docx.oxml.OxmlElement('w:rFonts')
            rPr.append(rFonts)
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)

def populate_table_0(table, items, total_amount, baht_text_str):
    num_items = len(items)
    total_rows = len(table.rows)
    total_row_idx = total_rows - 1
    template_item_rows_count = total_rows - 3
    
    if num_items < template_item_rows_count:
        for r_idx in range(total_row_idx - 1, 1 + num_items, -1):
            tr = table.rows[r_idx]._tr
            tr.getparent().remove(tr)
    elif num_items > template_item_rows_count:
        sample_tr = table.rows[2]._tr
        last_tr = table.rows[-1]._tr
        for _ in range(num_items - template_item_rows_count):
            new_tr = copy.deepcopy(sample_tr)
            last_tr.addprevious(new_tr)
            
    for idx, item in enumerate(items):
        row = table.rows[2 + idx]
        qty_unit = f"{item.get('qty', '')} {item.get('unit', '')}".strip()
        price_b, price_s = split_baht_satang(item.get('price_per_unit', 0))
        tot_b, tot_s = split_baht_satang(item.get('total_price', 0))
        
        set_cell_formatted_text(row.cells[0], str(idx + 1), WD_ALIGN_PARAGRAPH.CENTER, bold=False)
        set_cell_formatted_text(row.cells[1], str(item.get('name', '')), WD_ALIGN_PARAGRAPH.LEFT, bold=False)
        set_cell_formatted_text(row.cells[2], qty_unit, WD_ALIGN_PARAGRAPH.CENTER, bold=False)
        set_cell_formatted_text(row.cells[3], price_b, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
        set_cell_formatted_text(row.cells[4], price_s, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
        set_cell_formatted_text(row.cells[5], price_b, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
        set_cell_formatted_text(row.cells[6], price_s, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
        set_cell_formatted_text(row.cells[7], tot_b, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
        set_cell_formatted_text(row.cells[8], tot_s, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
        
    last_row = table.rows[-1]
    tot_b, tot_s = split_baht_satang(total_amount)
    set_cell_formatted_text(last_row.cells[0], f"รวมเป็นเงินทั้งสิ้น  ({baht_text_str})", WD_ALIGN_PARAGRAPH.CENTER, bold=False)
    set_cell_formatted_text(last_row.cells[7], tot_b, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
    set_cell_formatted_text(last_row.cells[8], tot_s, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)

def populate_table_1(table, items, financial, baht_text_str):
    num_items = len(items)
    total_rows = len(table.rows)
    template_item_rows_count = total_rows - 4
    
    if num_items < template_item_rows_count:
        for r_idx in range(total_rows - 4, num_items, -1):
            tr = table.rows[r_idx]._tr
            tr.getparent().remove(tr)
    elif num_items > template_item_rows_count:
        sample_tr = table.rows[1]._tr
        first_summary_tr = table.rows[-3]._tr
        for _ in range(num_items - template_item_rows_count):
            new_tr = copy.deepcopy(sample_tr)
            first_summary_tr.addprevious(new_tr)
            
    for idx, item in enumerate(items):
        row = table.rows[1 + idx]
        p_unit_str = format_money(item.get('price_per_unit', 0))
        tot_str = format_money(item.get('total_price', 0))
        
        set_cell_formatted_text(row.cells[0], str(idx + 1), WD_ALIGN_PARAGRAPH.CENTER, bold=False)
        set_cell_formatted_text(row.cells[1], str(item.get('name', '')), WD_ALIGN_PARAGRAPH.LEFT, bold=False)
        set_cell_formatted_text(row.cells[2], str(item.get('qty', '')), WD_ALIGN_PARAGRAPH.CENTER, bold=False)
        set_cell_formatted_text(row.cells[3], str(item.get('unit', '')), WD_ALIGN_PARAGRAPH.CENTER, bold=False)
        set_cell_formatted_text(row.cells[4], p_unit_str, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
        set_cell_formatted_text(row.cells[5], tot_str, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
        
    subtotal_row = table.rows[-3]
    vat_row = table.rows[-2]
    total_row = table.rows[-1]
    
    # Set BahtText ONCE in the merged bottom-left cell (clean single paragraph, vertically centered)
    bottom_left_cell = subtotal_row.cells[0]
    set_cell_formatted_text(bottom_left_cell, f"({baht_text_str})", align=WD_ALIGN_PARAGRAPH.CENTER, bold=False)
    bottom_left_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        
    subtotal_val = format_money(financial.get('goods_value', financial.get('total_amount', 0)))
    vat_val = format_money(financial.get('vat', 0)) if financial.get('has_vat') else "-"
    total_val = format_money(financial.get('total_amount', 0))
    
    set_cell_formatted_text(subtotal_row.cells[5], subtotal_val, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
    set_cell_formatted_text(vat_row.cells[5], vat_val, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)
    set_cell_formatted_text(total_row.cells[5], total_val, WD_ALIGN_PARAGRAPH.RIGHT, bold=False)

def apply_tag_replacements(doc, mapping: dict):
    sorted_tags = sorted(mapping.keys(), key=len, reverse=True)
    for p in doc.paragraphs:
        if '{{' in p.text:
            for tag in sorted_tags:
                val = mapping[tag]
                while tag in p.text:
                    pos = replace_text_in_paragraph(p, tag, val)
                    if pos == -1:
                        break
                        
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if '{{' in p.text:
                        for tag in sorted_tags:
                            val = mapping[tag]
                            while tag in p.text:
                                pos = replace_text_in_paragraph(p, tag, val)
                                if pos == -1:
                                    break

def generate_procurement_doc(data: dict, output_path: str) -> str:
    doc_type = data.get("doc_type", "buy")
    approver_type = data.get("approver_type", "director")
    
    template_file = get_template_path(doc_type, approver_type)
    if not os.path.exists(template_file):
        raise FileNotFoundError(f"Template not found: {template_file}")
        
    doc = docx.Document(template_file)
    
    dept = str(data.get("department", "ฝ่ายบริหารทั่วไป")).strip()
    project_name = str(data.get("project_name", "")).strip()
    hire_job_name = str(data.get("hire_job_name", "ทำป้ายไวนิล")).strip()
    items = data.get("items", [])
    num_items = len(items)
    
    total_amount = sum(float(item.get("total_price", 0)) for item in items)
    has_vat = data.get("has_vat", False)
    if has_vat:
        goods_value = round(total_amount * 100 / 107, 2)
        vat = round(total_amount - goods_value, 2)
    else:
        goods_value = total_amount
        vat = 0.0
        
    has_wht = data.get("has_withholding_tax", False)
    wht_rate = float(data.get("withholding_rate", 1.0)) if has_wht else 0.0
    withholding_tax = round(goods_value * (wht_rate / 100), 2) if has_wht else 0.0
    fine = float(data.get("fine_amount", 0.0)) if data.get("fine_amount") else 0.0
    net_pay = round(total_amount - withholding_tax - fine, 2)
    
    baht_text_total = baht_text(total_amount)
    baht_text_net = baht_text(net_pay)
    
    financial = {
        "goods_value": goods_value,
        "has_vat": has_vat,
        "vat": vat,
        "total_amount": total_amount,
        "has_withholding_tax": has_wht,
        "withholding_tax": withholding_tax,
        "fine": fine,
        "net_pay": net_pay
    }
    
    doc1_no = clean_doc_no(data.get("doc1_no", ""))
    doc1_date = str(data.get("doc1_date", "")).strip()
    doc2_no = clean_doc_no(data.get("doc2_no", ""))
    doc2_date = str(data.get("doc2_date", "")).strip()
    doc3_no = clean_doc_no(data.get("doc3_no", ""))
    doc3_date = str(data.get("doc3_date", "")).strip()
    
    year_req = extract_thai_year(doc1_date)
    
    po_no = str(data.get("po_no", "")).strip()
    po_date = str(data.get("po_date", "")).strip()
    delivery_days = str(data.get("delivery_days", 5)).strip()
    delivery_due_date = str(data.get("delivery_due_date", "")).strip()
    delivery_actual_date = str(data.get("delivery_actual_date", "")).strip()
    
    vendor = data.get("vendor", {})
    vendor_name = str(vendor.get("name", "")).strip()
    vendor_tax = str(vendor.get("tax_id", "")).strip()
    vendor_raw_addr = str(vendor.get("address", "")).strip()
    vendor_house, vendor_moo = parse_vendor_address(vendor_raw_addr)
    if vendor.get("moo"):
        vendor_moo = str(vendor.get("moo")).strip()
        
    vendor_subdist = str(vendor.get("subdistrict", "")).strip()
    vendor_dist = str(vendor.get("district", "")).strip()
    vendor_prov = str(vendor.get("province", "")).strip()
    vendor_zip = str(vendor.get("zipcode", "")).strip()
    vendor_phone = str(vendor.get("phone", "-")).strip()
    vendor_signer = str(vendor.get("signer_name", "")).strip()
    vendor_signer_pos = str(vendor.get("signer_position", "เจ้าของกิจการ")).strip()
    
    receipt = data.get("receipt", {})
    receipt_type = str(receipt.get("type", "ใบเสร็จรับเงิน")).strip()
    receipt_book = str(receipt.get("book_no", "-")).strip()
    receipt_no = str(receipt.get("no", "-")).strip()
    
    committee = data.get("committee", [
        {"name": "นางสาวธัญญาภรณ์  สุกันทา", "position": "ครู"},
        {"name": "นางสาวกนกวรรณ  มีเทียม", "position": "ครู"},
        {"name": "นายภูวดล  บุญต่อ", "position": "ครูผู้ช่วย"}
    ])
    while len(committee) < 3:
        committee.append({"name": "-", "position": "-"})
        
    officers = data.get("officers") or {}
    officer_supplies = clean_officer_name(officers.get("officer_supplies", "นางสาวกรรณิกา  พึ่งทอง"))
    head_supplies = clean_officer_name(officers.get("head_supplies", "นางสาวธิดาภรณ์  คงชนะ"))
    finance_officer = clean_officer_name(officers.get("finance_officer", "นายสุรพล  คงยืน"))
    
    if len(doc.tables) > 0:
        populate_table_0(doc.tables[0], items, total_amount, baht_text_total)
        
    if len(doc.tables) > 1:
        populate_table_1(doc.tables[1], items, financial, baht_text_total)
        
    tot_b, tot_s = split_baht_satang(total_amount)
    goods_b, goods_s = split_baht_satang(goods_value)
    net_b, net_s = split_baht_satang(net_pay)
    
    mapping = {
        # 1. Master Template Actual Tags (Exact matches in Form_*.docx)
        "{{ปีที่ขอ}}": year_req,
        "{{ฝ่ายงาน/ผู้ขอ}}": dept,
        "{{เหตุผลความจำเป็น}}": project_name if doc_type == "buy" else hire_job_name,
        "{{จำนวนของที่ซื้อ}}": str(num_items),
        "{{จำนวนรายการจ้าง}}": str(num_items),
        "{{ชื่องานที่จะจ้าง}}": hire_job_name,
        "{{ราคา}}": tot_b,
        "{{ราคา.00}}": format_money(total_amount),
        "{{ราคาจริง.00}}": format_money(goods_value),
        "{{ภาษี.00}}": format_money(vat) if has_vat else "-",
        "{{หักภาษี.00}}": format_money(withholding_tax) if has_wht else "-",
        "{{ค่าปรับ.00}}": format_money(fine) if fine > 0 else "-",
        "{{ราคาจ่ายจริง.00}}": format_money(net_pay),
        "{{Bath text}}": baht_text_total,
        "{{Bath text จ่ายจริง}}": baht_text_net,
        "{{กำหนดส่ง}}": delivery_days,
        "{{วันครบกำหนดส่งมอบตามใบสั่ง}}": delivery_due_date,
        "{{วันที่บันทึกรายงานขอ}}": doc1_date,
        "{{เลขที่1}}": doc1_no,
        "{{วันที่อนุมัติสั่ง}}": doc2_date,
        "{{เลขที่2}}": doc2_no,
        "{{วันที่ใบสั่งซื้อ}}": po_date,
        "{{เลขที่ใบสั่งซื้อ}}": po_no,
        "{{วันที่ใบสั่งจ้าง}}": po_date,
        "{{เลขที่ใบสั่งจ้าง}}": po_no,
        "{{วันที่ตรวจรับงาน}}": delivery_actual_date,
        "{{วันที่บันทึกขออนุมัติจ่ายเงิน}}": doc3_date,
        "{{เลขที่3}}": doc3_no,
        "{{ชื่อร้านค้า/บริษัท}}": vendor_name,
        "{{เลขประจำตัวผู้เสียภาษี}}": vendor_tax,
        "{{เลขที่ร้าน}}": vendor_house,
        "{{หมู่ร้าน}}": vendor_moo,
        "{{ตำบล}}": vendor_subdist,
        "{{อำเภอ}}": vendor_dist,
        "{{จังหวัด}}": vendor_prov,
        "{{เลขไปรษณีย์}}": vendor_zip,
        "{{โทรศัพท์}}": vendor_phone,
        "{{ชื่อผู้ลงนามของผู้ขาย}}": vendor_signer,
        "{{ตำแหน่งของผู้ลงนาม}}": vendor_signer_pos,
        "{{ประเภทใบเสร็จ}}": receipt_type,
        "{{เล่มที่}}": receipt_book,
        "{{เลขที่}}": receipt_no,
        "{{ประธานกรรมการ}}": committee[0]["name"],
        "{{ประธาน 1}}": committee[0]["position"],
        "{{กรรมการคนที่ 1}}": committee[1]["name"],
        "{{กก 1}}": committee[1]["position"],
        "{{กรรมการคนที่ 2}}": committee[2]["name"],
        "{{กก 2}}": committee[2]["position"],
        "{{เจ้าหน้าที่พัสดุ}}": officer_supplies,
        "{{หัวหน้าเจ้าหน้าที่พัสดุ}}": head_supplies,
        "{{เจ้าหน้าที่การเงิน}}": finance_officer,
        
        # 2. Compatibility aliases
        "{{กลุ่มงาน / ฝ่าย / หน่วยบริการผู้ขอ}}": dept,
        "{{กลุ่มงาน / ฝ่าย / หน่วยบริการ}}": dept,
        "{{กลุ่มงาน / ฝ่าย}}": dept,
        "{{กลุ่มงาน/ฝ่าย}}": dept,
        "{{ชื่องานซื้อ}}": project_name,
        "{{ชื่องานจ้าง}}": hire_job_name,
        "{{รายการพัสดุ}}": project_name if doc_type == "buy" else hire_job_name,
        "{{จำนวนรายการ}}": str(num_items),
        "{{จำนวนเงินรวม}}": format_money(total_amount),
        "{{ราคารวม}}": format_money(total_amount),
        "{{ราคากลาง}}": format_money(total_amount),
        "{{ราคาบาท}}": tot_b,
        "{{ราคาสตางค์}}": tot_s,
        "{{bath text}}": baht_text_total,
        "{{ตัวหนังสือ}}": baht_text_total,
        "{{ภาษีมูลค่าเพิ่ม}}": format_money(vat) if has_vat else "-",
        "{{ภาษี 1%}}": format_money(withholding_tax) if has_wht else "-",
        "{{ค่าปรับ}}": format_money(fine) if fine > 0 else "-",
        "{{จำนวนเงินหักภาษี}}": format_money(net_pay),
        "{{สุทธิบาท}}": net_b,
        "{{สุทธิต่าง}}": net_s,
        "{{สุทธิต่างค์}}": net_s,
        "{{Bath text สุทธิ}}": baht_text_net,
        "{{เลขที่ 1}}": doc1_no,
        "{{เลขที่ 2}}": doc2_no,
        "{{เลขที่ 3}}": doc3_no,
        "{{ปีงบ}}": year_req,
        "{{วันที่สั่งซื้อ}}": po_date,
        "{{วันที่สั่งจ้าง}}": po_date,
        "{{กำหนดส่งมอบวัน}}": delivery_days,
        "{{วันที่ครบกำหนด}}": delivery_due_date,
        "{{วันที่ส่งมอบของจริง}}": delivery_actual_date,
        "{{ชื่อร้านค้า/ผู้รับจ้าง/ผู้ขาย}}": vendor_name,
        "{{ชื่อร้านค้า}}": vendor_name,
        "{{เลขผู้เสียภาษี}}": vendor_tax,
        "{{บ้านเลขที่}}": vendor_house,
        "{{หมู่}}": vendor_moo,
        "{{วันที่รายงานผล}}": doc2_date,
        "{{วันที่บันทึกตรวจรับ}}": doc3_date,
    }
    
    apply_tag_replacements(doc, mapping)
    
    # --- Disable proofing (spelling/grammar lines) dynamically ---
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    
    # 1. Add <w:noProof/> to every run in the document
    for p in doc.paragraphs:
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            if rPr.find(qn('w:noProof')) is None:
                rPr.append(OxmlElement('w:noProof'))
                
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        rPr = r._r.get_or_add_rPr()
                        if rPr.find(qn('w:noProof')) is None:
                            rPr.append(OxmlElement('w:noProof'))

    # 2. Add document-level settings to hide errors
    settings = doc.settings.element
    for tag in ['w:hideSpellingErrors', 'w:hideGrammaticalErrors']:
        if settings.find(qn(tag)) is None:
            settings.append(OxmlElement(tag))
            
    proofState = settings.find(qn('w:proofState'))
    if proofState is None:
        proofState = OxmlElement('w:proofState')
        settings.append(proofState)
    proofState.set(qn('w:spelling'), 'clean')
    proofState.set(qn('w:grammar'), 'clean')
    # -----------------------------------------------------------
    
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path
